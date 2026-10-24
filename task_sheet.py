"""Portable, consistent task briefings and bounded local download publication."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
import io
import math
from pathlib import Path
import secrets
from urllib.parse import urlsplit
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from config import LOCAL_TIMEZONE, METRES_TO_FEET, MPS_TO_KNOTS
from soaring_physics import distance_bearing, wind_direction
from task_planning import export_cup, export_tsk, task_distance, validate_task
from briefing_bot import format_windows

FONT_PATH = Path(__file__).parent/'assets'/'fonts'/'DejaVuSans.ttf'

NOTES = ('Forecast planning estimate. Check current airspace, NOTAMs, weather and the aircraft flight manual.',
         'Observation zones: 5 km start line; BGA fixed-course 20 km sector plus 0.5 km barrel; 1 km finish ring.',
         'These are declared planning defaults. Confirm the rules and imported zones on your instrument.',
         'Airfield candidates are published waypoints, not verified emergency landing sites. Check permission, runway and serviceability.',
         'Wind, heading and speed are leg averages. Conditions change between thermals. Magnetic tracks use WMM2025.')


def dm(value, latitude=True):
    minutes=round(abs(value)*60,3); degrees=int(minutes//60)
    return f'{degrees:0{2 if latitude else 3}d}° {minutes-degrees*60:06.3f}′ '+(('N' if value>=0 else 'S') if latitude else ('E' if value>=0 else 'W'))


def duration(hours):
    if hours is None: return 'Unavailable'
    minutes=round(hours*60)
    return f'{minutes//60:02d}:{minutes%60:02d}'


def magnetic_variation(point, day):
    from pygeomag import GeoMag
    beginning=datetime(day.year,1,1); following=datetime(day.year+1,1,1)
    year=day.year+(datetime.combine(day,datetime.min.time())-beginning).total_seconds()/(following-beginning).total_seconds()
    try: return GeoMag().calculate(glat=point.latitude,glon=point.longitude,alt=(point.elevation_m or 0)/1000,time=year).d
    except ValueError: return None  # Never extrapolate a magnetic model beyond its validity.


def wind_heading(track, airspeed_kmh, u, v):
    if airspeed_kmh is None or airspeed_kmh<=0: return None
    angle=math.radians(track); cross=u*math.cos(angle)-v*math.sin(angle)
    ratio=cross/(airspeed_kmh/3.6)
    if abs(ratio)>=1: return None
    return (track-math.degrees(math.asin(ratio)))%360


def nearby_airfields(points, catalog):
    """Bounded distance to each leg in a local metric projection, not suitability."""
    result=[]
    for a,b in zip(points,points[1:]):
        scale=111*math.cos(math.radians((a.latitude+b.latitude)/2))
        dx=(b.longitude-a.longitude)*scale;dy=(b.latitude-a.latitude)*111
        hits=[]
        for p in catalog:
            if p.style not in (2,3,4,5): continue
            px=(p.longitude-a.longitude)*scale;py=(p.latitude-a.latitude)*111
            t=max(0,min(1,(px*dx+py*dy)/(dx*dx+dy*dy)))
            distance=math.hypot(px-t*dx,py-t*dy)
            if distance<=15: hits.append((distance,p))
        hits.sort(key=lambda pair:pair[0])
        details=[]
        for distance,p in hits[:4]:
            runway='' if p.runway_length_m is None else f'; published runway {p.runway_length_m:.0f} m'
            details.append(f'{p.code} {p.name}: {distance:.0f} km from leg{runway}. Verify availability.')
        result.append(f'{a.code} → {b.code}: '+(' '.join(details) if details else 'No published airfield candidate within 15 km. Plan landing options separately.'))
    return tuple(result)


@dataclass(frozen=True)
class TaskSheet:
    header: dict
    rows: tuple[dict,...]
    airfields: tuple[str,...]
    notes: tuple[str,...]


def build_sheet(points, assessment, day, loading, catalog=()):
    points=validate_task(points); best=assessment.best if assessment else None
    distance=task_distance(points)
    header={'Aircraft':'G-CKML · Schempp-Hirth Duo Discus','Date':day.isoformat(),
            'Start window':format_windows(assessment)+' Europe/London' if assessment else 'Unavailable',
            'Wing loading':'34 kg/m² · dry' if loading=='dry' else '43 kg/m² · water ballast',
            'Task distance':f'{distance:.1f} km / {distance/1.852:.1f} nm',
            'Distance class':next((f'{k} km' for k in (500,300,250,200,100) if distance>=k),'Local task'),
            'Verdict':assessment.verdict if assessment else 'INSUFFICIENT DATA',
            'Predicted XC speed':'Unavailable' if not best else f'{best.average_groundspeed_kmh:.1f} km/h / {best.average_groundspeed_kmh/1.852:.1f} kt',
            'Total elapsed':duration(best.elapsed_hours if best else None)}
    rows=[]
    legs={leg.leg_index:leg for leg in best.legs} if best else {}
    for i,p in enumerate(points):
        role='Start' if i==0 else 'Finish' if i==len(points)-1 else f'TP{i}'
        zone='Start line 5 km' if i==0 else 'Finish ring 1 km' if i==len(points)-1 else 'BGA sector 20 km / barrel 0.5 km'
        row={'Sequence':role,'Code':p.code,'Turnpoint':p.name,'Latitude':dm(p.latitude),
             'Longitude':dm(p.longitude,False),'Elevation ft AMSL':'Unknown' if p.elevation_m is None else f'{p.elevation_m*METRES_TO_FEET:.0f}',
             'Observation zone':zone, 'Leg km / nm':'—','Track true / magnetic':'—','Wind kt / from °':'—',
             'Heading true °':'—','STF kt':'—','XC km/h / kt':'—','Leg time hh:mm':'—'}
        if i:
            length,track=distance_bearing(points[i-1].point,p.point)
            variation=magnetic_variation(points[i-1],day)
            row['Leg km / nm']=f'{length:.1f} / {length/1.852:.1f}'
            row['Track true / magnetic']=f'{track:.0f} / '+('Unknown' if variation is None else f'{(track-variation)%360:.0f}')
            leg=legs.get(i-1)
            for key in ('Wind kt / from °','Heading true °','STF kt','XC km/h / kt','Leg time hh:mm'): row[key]='Unavailable'
            if leg:
                direction=wind_direction(leg.wind_u_mps,leg.wind_v_mps)
                row['Wind kt / from °']=f'{math.hypot(leg.wind_u_mps,leg.wind_v_mps)*MPS_TO_KNOTS:.1f} / '+('calm' if direction is None else f'{direction:.0f}')
                heading=wind_heading(track,leg.cruise_airspeed_kmh,leg.wind_u_mps,leg.wind_v_mps)
                row['Heading true °']='Unavailable' if heading is None else f'{heading:.0f}'
                row['STF kt']=f'{leg.cruise_airspeed_kmh/1.852:.1f}'
                row['XC km/h / kt']=f'{leg.groundspeed_kmh:.1f} / {leg.groundspeed_kmh/1.852:.1f}'
                row['Leg time hh:mm']=duration(leg.elapsed_hours)
        rows.append(row)
    return TaskSheet(header,tuple(rows),nearby_airfields(points,catalog),NOTES)


def map_snapshot(points):
    """Offline geographic route snapshot with land mask, labels and scale grid."""
    import numpy as np
    from global_land_mask import globe
    image=Image.new('RGB',(1200,650),'#e8f0f3');draw=ImageDraw.Draw(image)
    west=min(p.longitude for p in points)-.35;east=max(p.longitude for p in points)+.35
    south=min(p.latitude for p in points)-.25;north=max(p.latitude for p in points)+.25
    # Preserve local distance aspect ratio, centred on the task.
    mid=(south+north)/2; ratio=math.cos(math.radians(mid))
    span=max((east-west)*ratio/1120,(north-south)/530)
    cx=(east+west)/2;cy=mid
    west=cx-span*1120/ratio/2;east=cx+span*1120/ratio/2
    south=cy-span*530/2;north=cy+span*530/2
    xs=np.linspace(west,east,560);ys=np.linspace(north,south,265);xx,yy=np.meshgrid(xs,ys)
    land=globe.is_land(yy,xx);pixels=np.where(land[...,None],np.array([239,239,222]),np.array([225,237,244])).astype('uint8')
    image.paste(Image.fromarray(pixels).resize((1120,530)),(40,60));draw=ImageDraw.Draw(image)
    try:font=ImageFont.truetype(str(FONT_PATH),19);small=ImageFont.truetype(str(FONT_PATH),15)
    except OSError:font=small=ImageFont.load_default()
    def xy(p):return 40+(p.longitude-west)/(east-west)*1120,60+(north-p.latitude)/(north-south)*530
    for lat in np.arange(math.ceil(south*2)/2,north,.5):
        y=60+(north-lat)/(north-south)*530;draw.line((40,y,1160,y),fill='#ccd5d4');draw.text((45,y),f'{lat:.1f} N',font=small,fill='#566779')
    draw.line([xy(p) for p in points],fill='#234e76',width=5)
    for i,p in enumerate(points):
        x,y=xy(p);draw.ellipse((x-7,y-7,x+7,y+7),fill='#f5aa4d',outline='#253b4b',width=2)
        # Repeated start/finish gets one combined label.
        if i==len(points)-1 and p==points[0]: continue
        label=f'{i+1} {p.code}';draw.text((min(x+9,1050),max(65,y-25)),label,font=font,fill='#152d43')
    draw.text((40,15),'G-CKML · task route overview',font=font,fill='#1e242b')
    draw.text((40,610),'Geographic overview • no airspace, terrain heights or landing clearance shown',font=small,fill='#1e242b')
    buffer=io.BytesIO();image.save(buffer,format='PNG');return buffer.getvalue()


def pdf_bytes(sheet, snapshot):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font_path=FONT_PATH
    font='Helvetica'
    if font_path.exists():
        if 'DuoSans' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('DuoSans',str(font_path)))
        font='DuoSans'
    # Standard Helvetica fallback gets ASCII transliteration, never missing glyphs.
    def clean(text):
        text=str(text)
        return text if font=='DuoSans' else text.replace('→',' to ').replace('—','-').replace('′',"'").replace('·',' / ').replace('²','2').replace('–','-')
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='Cell',fontName=font,fontSize=7,leading=10))
    styles.add(ParagraphStyle(name='BodyDuo',fontName=font,fontSize=9,leading=14,spaceAfter=7))
    p=lambda s:Paragraph(escape(clean(s)),styles['Cell'])
    buf=io.BytesIO();doc=SimpleDocTemplate(buf,pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=28,bottomMargin=30)
    story=[Paragraph('BGA / SeeYou Task Sheet',styles['Title'])]
    header=Table([[p(k),p(v)] for k,v in sheet.header.items()],colWidths=[125,660])
    header.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#f4f6f9')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story.append(header)
    story.extend([Spacer(1,8),RLImage(io.BytesIO(snapshot),width=600,height=325),PageBreak()])
    groups=[['Sequence','Code','Turnpoint','Latitude','Longitude','Elevation ft AMSL','Observation zone','Leg km / nm','Track true / magnetic'],
            ['Sequence','Code','Wind kt / from °','Heading true °','STF kt','XC km/h / kt','Leg time hh:mm']]
    for columns in groups:
        data=[[p(c) for c in columns]]+[[p(row[c]) for c in columns] for row in sheet.rows]
        table=Table(data,repeatRows=1,colWidths=[785/len(columns)]*len(columns))
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dce6eb')),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#9aabb6')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        story.extend([table,Spacer(1,16)])
    for note in (*sheet.airfields,*sheet.notes):story.append(Paragraph(escape(clean(note)),styles['BodyDuo']))
    def footer(canvas,document):
        canvas.setFont(font,8);canvas.drawString(28,16,'G-CKML | Forecast task briefing | Check current flight information');canvas.drawRightString(812,16,str(document.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer);return buf.getvalue()


def docx_bytes(sheet,snapshot):
    from docx import Document
    from docx.shared import Inches, Pt, Mm
    from docx.enum.section import WD_ORIENT
    document=Document();section=document.sections[0];section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width=Mm(297);section.page_height=Mm(210)
    section.top_margin=section.bottom_margin=Mm(12);section.left_margin=section.right_margin=Mm(12)
    style=document.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10);style.paragraph_format.line_spacing=1.4
    document.add_heading('BGA / SeeYou Task Sheet',0)
    table=document.add_table(rows=0,cols=2);table.style='Light Shading Accent 1'
    for key,value in sheet.header.items():cells=table.add_row().cells;cells[0].text=key;cells[1].text=value
    document.add_picture(io.BytesIO(snapshot),width=Inches(8.5))
    document.add_page_break()
    columns=list(sheet.rows[0]);groups=(columns[:9],['Sequence','Code',*columns[9:]])
    for group in groups:
        table=document.add_table(rows=1,cols=len(group));table.style='Light Shading Accent 1'
        for cell,name in zip(table.rows[0].cells,group):cell.text=name
        # Repeat column headings on multi-page tables.
        from docx.oxml import OxmlElement
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in sheet.rows:
            for cell,key in zip(table.add_row().cells,group):cell.text=row[key]
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:run.font.size=Pt(8)
        document.add_paragraph()
    document.add_heading('Airfield candidates and briefing notes',1)
    for note in (*sheet.airfields,*sheet.notes):document.add_paragraph(note)
    out=io.BytesIO();document.save(out);return out.getvalue()


def export_bundle(points,assessment,day,loading,catalog=()):
    sheet=build_sheet(points,assessment,day,loading,catalog);snapshot=map_snapshot(points)
    return sheet, {'pdf':pdf_bytes(sheet,snapshot),'docx':docx_bytes(sheet,snapshot),
                   'cup':export_cup(points,f'G-CKML {day}',bga_zones=True).encode('utf-8'),
                   'tsk':export_tsk(points,bga_zones=True).encode('utf-8'),'png':snapshot}


def publish_downloads(bundle, *, root=None):
    """Publish only generated task files behind an unguessable, local static path."""
    root=Path(root or Path(__file__).parent/'static'/'tasks');root.mkdir(parents=True,exist_ok=True)
    token=secrets.token_urlsafe(24);folder=root/token;folder.mkdir()
    for extension,data in bundle.items():
        if extension not in ('pdf','docx','cup','tsk','png'):raise ValueError('Unsupported export format')
        (folder/f'duo-task.{extension}').write_bytes(data)
    # Task files expire after 30 days; never touch unrelated static content.
    import re,shutil,time
    for entry in root.iterdir():
        if entry.is_dir() and re.fullmatch(r'[A-Za-z0-9_-]{32}',entry.name) and time.time()-entry.stat().st_mtime>30*86400:
            shutil.rmtree(entry)
    return f'app/static/tasks/{token}/duo-task'


def download_url(base,path,extension):
    parsed=urlsplit(base)
    if parsed.scheme not in ('https','http') or not parsed.netloc or parsed.username or parsed.password or parsed.query or parsed.fragment:
        raise ValueError('Use the application’s HTTP(S) address, without a query or fragment')
    if extension not in ('cup','tsk','pdf','docx'):raise ValueError('Unsupported format')
    return base.rstrip('/')+'/'+path+'.'+extension


def qr_bytes(url):
    import qrcode
    out=io.BytesIO();qrcode.make(url).save(out,format='PNG');return out.getvalue()
