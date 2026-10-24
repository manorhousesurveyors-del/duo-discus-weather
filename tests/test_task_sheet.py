from dataclasses import replace
from datetime import date
import csv
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

from PIL import Image
from docx import Document
from tests.test_task_planning import TASK,weather,DAY
from task_planning import assess_task,export_cup,export_tsk,candidate_tasks,task_distance
from task_sheet import (build_sheet,export_bundle,wind_heading,dm,magnetic_variation,
                        nearby_airfields,publish_downloads,download_url,qr_bytes)
from bga_turnpoints import load_catalog,parse_turnpoints
from route_tools import choose_optimised


class TaskSheetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.assessment=assess_task(TASK,weather(),DAY)
        cls.sheet,cls.files=export_bundle(TASK,cls.assessment,DAY,'dry')

    def test_leg_totals_match_the_integrated_flight_not_hourly_snapshot(self):
        best=self.assessment.best
        self.assertEqual(len(best.legs),len(TASK)-1)
        self.assertAlmostEqual(sum(l.distance_km for l in best.legs),self.assessment.distance_km)
        self.assertAlmostEqual(sum(l.elapsed_hours for l in best.legs),best.elapsed_hours)
        self.assertEqual(best.legs[0].finished_at,best.legs[1].started_at)
        self.assertNotEqual(best.legs[0].groundspeed_kmh,best.legs[1].groundspeed_kmh)
        self.assertTrue(all(float(row['STF kt'])>50 for row in self.sheet.rows[1:]))

    def test_pdf_word_and_map_are_valid_containers_with_complete_tables(self):
        pdf=self.files['pdf'];self.assertTrue(pdf.startswith(b'%PDF-'));self.assertIn(b'%%EOF',pdf)
        self.assertIn(b'/Subtype /Image',pdf);self.assertIn(b'/Type /Pages',pdf)
        document=Document(io.BytesIO(self.files['docx']))
        text=' '.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        for value in ('G-CKML','FULLY ACHIEVABLE','Start line 5 km','BGA sector 20 km / barrel 0.5 km','Heading true','STF kt'):
            self.assertIn(value,text)
        self.assertEqual(len(document.tables[1].rows),len(TASK)+1)
        self.assertEqual(Image.open(io.BytesIO(self.files['png'])).size,(1200,650))

    def test_cup_task_definition_and_observation_zones_match_tsk(self):
        cup=self.files['cup'].decode();self.assertEqual(len(parse_turnpoints(cup)),2)
        rows=list(csv.reader(io.StringIO(cup)))
        taskrow=rows[next(i for i,r in enumerate(rows) if r[0]=='-----Related Tasks-----')+1]
        self.assertEqual(taskrow[1:],[TASK[0].name,TASK[0].name,TASK[1].name,TASK[0].name,TASK[0].name])
        self.assertIn('ObsZone=0,Style=2,R1=2500m,A1=180,Line=1',cup)
        self.assertIn('R1=20000m,A1=45,R2=500m,A2=180',cup)
        root=ET.fromstring(self.files['tsk'])
        zones=[point.find('ObservationZone').attrib for point in root.findall('Point')]
        self.assertEqual(zones,[{'type':'Line','length':'5000'},{'type':'BGAFixedCourse'},{'type':'Cylinder','radius':'1000'}])

    def test_wind_correction_and_magnetic_track_have_correct_sign_and_model_limits(self):
        self.assertAlmostEqual(wind_heading(0,100,0,10),0)
        self.assertGreater(wind_heading(0,100,10,0),330) # wind pushes east; point west of north
        self.assertIsNone(wind_heading(0,10,10,0))
        self.assertAlmostEqual(magnetic_variation(TASK[0],date(2026,10,24)),.22,delta=.5)
        self.assertIsNone(magnetic_variation(TASK[0],date(2040,1,1)))
        self.assertEqual(dm(52+59.9999/60),'53° 00.000′ N')

    def test_unachievable_task_exports_without_inventing_flight_numbers(self):
        result=assess_task(TASK,weather(precipitation=2),DAY)
        sheet=build_sheet(TASK,result,DAY,'ballasted')
        self.assertEqual(sheet.header['Predicted XC speed'],'Unavailable')
        self.assertEqual(sheet.rows[1]['STF kt'],'Unavailable')
        self.assertIn('43',sheet.header['Wing loading'])
        sheet=build_sheet(TASK,None,DAY,'dry')
        self.assertEqual(sheet.header['Verdict'],'INSUFFICIENT DATA')

    def test_published_downloads_use_unique_local_paths_and_verified_qr_target(self):
        with tempfile.TemporaryDirectory() as folder:
            a=publish_downloads(self.files,root=folder);b=publish_downloads(self.files,root=folder)
            self.assertNotEqual(a,b)
            path=Path(folder)/a.split('/')[-2]/'duo-task.tsk'
            self.assertEqual(path.read_bytes(),self.files['tsk'])
            url=download_url('https://weather.example/soaring/',a,'tsk')
            self.assertTrue(url.startswith('https://weather.example/soaring/app/static/tasks/'))
            qr=qr_bytes(url);self.assertGreater(Image.open(io.BytesIO(qr)).width,200)
        for base in ('javascript:alert(1)','https://user:password@host','https://host/?token=x'):
            with self.assertRaises(ValueError):download_url(base,'path','cup')

    def test_airfield_candidates_are_not_claimed_as_verified_outlandings(self):
        field=replace(TASK[0],style=2,code='AIR',runway_length_m=500)
        notes=nearby_airfields(TASK,(field,))
        self.assertIn('AIR',notes[0]);self.assertIn('Verify availability',notes[0])
        self.assertIn('No published airfield',nearby_airfields(TASK,())[0])


class OptimizerTests(unittest.TestCase):
    def test_new_200km_and_custom_targets_are_closed_bga_routes(self):
        catalog=load_catalog()
        for target in (200,275,500):
            candidates=candidate_tasks(catalog.turnpoints,catalog.by_code['MYN'],target)
            self.assertTrue(candidates)
            for task in candidates:
                self.assertEqual(task[0],task[-1]);self.assertTrue(target<=task_distance(task)<=target*1.08)
        with self.assertRaises(ValueError):candidate_tasks((),TASK[0],-1)

    def test_optimiser_excludes_failed_routes_and_uses_objective(self):
        good=assess_task(TASK,weather(),DAY);bad=assess_task(TASK,(),DAY)
        self.assertIsNone(choose_optimised([(TASK,bad)],'Fastest speed'))
        longer=replace(good,distance_km=300,best=replace(good.best,average_groundspeed_kmh=70))
        self.assertEqual(choose_optimised([(TASK,good),(TASK,longer)],'Fastest speed')[1],good)
        self.assertEqual(choose_optimised([(TASK,good),(TASK,longer)],'Best distance')[1],longer)
